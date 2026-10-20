from __future__ import annotations

import json
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from company_discovery import DiscoveryResult, extract_connector_candidates
from connectors import (
    _SALESFORCE_DETAIL_CACHE,
    ConnectorError,
    JobDetail,
    JobPosting,
    fetch_adobe,
    fetch_apple,
    fetch_job_detail,
    fetch_jobs,
    fetch_lever,
    fetch_meta,
    fetch_microsoft,
    fetch_salesforce,
    fetch_smartrecruiters,
    fetch_tesla,
)
from initialize_companies import company_key_from_name, initialize_companies
from job_finder import (
    _deterministic_rejection_reason,
    evaluate_retryable_postings,
    process_retryable_postings,
)
from llm_evaluator import EvaluationConfig, EvaluationError, evaluate_job
from policy import (
    clean_job_description,
    employment_type_prefilter,
    location_is_allowed,
    required_minimum_years,
    title_prefilter,
)
from storage import Store, make_signature
from workbook_io import (
    _append_table_row,
    append_run_rows,
    append_review_queue_jobs,
    read_companies,
    review_queue_contains,
)


POLICY = json.loads(
    (ROOT / "config" / "search_policy.example.json").read_text(encoding="utf-8")
)
PROFILE = json.loads(
    (ROOT / "config" / "candidate_profile.example.json").read_text(encoding="utf-8")
)


def evaluation_result(**overrides):
    result = {
        "minimum_years_required": 2,
        "experience_evidence": "Two years of relevant experience required.",
        "experience_pass": True,
        "sponsorship_status": "not_mentioned",
        "sponsorship_evidence": None,
        "sponsorship_pass": True,
        "role_fit": "high",
        "relevance_score": 88,
        "recommended_accept": True,
        "decision_reason": "Strong data engineering match.",
        "matching_evidence": ["Python and Spark", "Data platform ownership"],
        "gaps": [],
    }
    result.update(overrides)
    return result


class LlmEvaluatorTests(unittest.TestCase):
    def test_responses_request_uses_strict_schema_and_parses_output(self):
        captured = {}

        def transport(url, payload, api_key, timeout):
            captured.update(
                {"url": url, "payload": payload, "api_key": api_key, "timeout": timeout}
            )
            return {
                "status": "completed",
                "output": [
                    {
                        "type": "message",
                        "content": [
                            {
                                "type": "output_text",
                                "text": json.dumps(evaluation_result()),
                            }
                        ],
                    }
                ],
                "usage": {"input_tokens": 500, "output_tokens": 100, "total_tokens": 600},
            }

        config = EvaluationConfig(api_key="test-key", model="test-model")
        result = evaluate_job(
            job={"company": "Example", "title": "Data Engineer"},
            description="Build reliable data systems.",
            candidate_profile=PROFILE,
            search_policy=POLICY,
            prompt_directory=ROOT / "config",
            config=config,
            transport=transport,
        )

        self.assertEqual(result["relevance_score"], 88)
        self.assertEqual(result["api_usage"]["total_tokens"], 600)
        self.assertFalse(captured["payload"]["store"])
        self.assertTrue(captured["payload"]["text"]["format"]["strict"])
        self.assertEqual(captured["payload"]["reasoning"]["effort"], "medium")

    def test_hard_rules_override_model_recommendation(self):
        def transport(_url, _payload, _api_key, _timeout):
            return {
                "status": "completed",
                "output": [
                    {
                        "type": "message",
                        "content": [
                            {
                                "type": "output_text",
                                "text": json.dumps(
                                    evaluation_result(
                                        minimum_years_required=4,
                                        experience_pass=True,
                                        recommended_accept=True,
                                    )
                                ),
                            }
                        ],
                    }
                ],
            }

        result = evaluate_job(
            job={"company": "Example", "title": "Data Engineer"},
            description="Four years required.",
            candidate_profile=PROFILE,
            search_policy=POLICY,
            prompt_directory=ROOT / "config",
            config=EvaluationConfig(api_key="test-key"),
            transport=transport,
        )
        self.assertFalse(result["experience_pass"])
        self.assertFalse(result["recommended_accept"])


class PolicyTests(unittest.TestCase):
    def test_locations(self):
        allowed = [
            ("San Mateo, CA", "Hybrid"),
            ("Buffalo, NY", "On-site"),
            ("New York City", "Hybrid"),
            ("Boston, Massachusetts", "On-site"),
        ]
        for location, location_type in allowed:
            self.assertTrue(location_is_allowed(location, location_type, POLICY))
        self.assertFalse(location_is_allowed("United States", "Remote", POLICY))
        self.assertFalse(location_is_allowed("California", "Remote", POLICY))

    def test_title_rules(self):
        accepted = [
            "Data Engineer",
            "Machine Learning Engineer",
            "Product Manager, Data Platform",
            "Software Engineer, Data Infrastructure",
            "Forward Deployed Engineer",
        ]
        for title in accepted:
            self.assertTrue(title_prefilter(title, POLICY)[0], title)
        rejected = [
            "Software Engineer, Payments",
            "Staff Data Engineer",
            "Research Scientist",
            "Sales Engineer",
            "Data Engineering Intern",
            "Security GRC Lead, Governance Engineering",
        ]
        for title in rejected:
            self.assertFalse(title_prefilter(title, POLICY)[0], title)

    def test_remote_title_is_rejected_even_with_an_allowed_city(self):
        posting = JobPosting(
            "1",
            "Software Engineering SMTS / LMTS - AI Core (REMOTE)",
            "California - San Francisco",
            "https://example.com/jobs/1",
        )
        self.assertIn(
            "work arrangement",
            _deterministic_rejection_reason(posting, POLICY),
        )

    def test_experience_parser(self):
        self.assertEqual(required_minimum_years("At least 4 years of experience"), 4)
        self.assertEqual(required_minimum_years("3-5 years of experience"), 3)
        self.assertIsNone(required_minimum_years("Relevant experience preferred"))

    def test_employment_type_prefilter(self):
        self.assertTrue(employment_type_prefilter("Full-time", POLICY)[0])
        self.assertFalse(employment_type_prefilter("Contract", POLICY)[0])
        self.assertTrue(employment_type_prefilter("", POLICY)[0])

    def test_description_cleaner_removes_html_and_scripts(self):
        description = "<h2>Role</h2><p>Build&nbsp;data systems.</p><script>ignore()</script><li>Python</li>"
        self.assertEqual(
            clean_job_description(description),
            "Role\nBuild data systems.\nPython",
        )


class StorageTests(unittest.TestCase):
    def test_baseline_and_delta(self):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / "test.db")
            company = {
                "company_key": "example",
                "company": "Example",
                "connector": "greenhouse",
                "connector_key": "example",
                "careers_url": "https://example.com/careers",
            }
            store.sync_companies([company])
            run_1 = store.start_run("baseline")
            first = JobPosting("1", "Data Engineer", "New York, NY", "https://example.com/1")
            self.assertEqual(store.store_postings("example", [first], run_1, baseline=True), [])
            run_2 = store.start_run("run")
            second = JobPosting("2", "AI Engineer", "California", "https://example.com/2")
            new_jobs = store.store_postings("example", [first, second], run_2, baseline=False)
            self.assertEqual([job.source_job_id for job in new_jobs], ["2"])
            self.assertIsNone(
                store.get_processing_state("example", make_signature("example", "1"))
            )
            state = store.get_processing_state(
                "example", make_signature("example", "2")
            )
            self.assertEqual(state["status"], "discovered")
            store.close()

    def test_processing_failures_remain_retryable_until_review_queue(self):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / "test.db")
            company = {
                "company_key": "example",
                "company": "Example",
                "connector": "greenhouse",
                "connector_key": "example",
                "careers_url": "https://example.com/careers",
            }
            store.sync_companies([company])
            run_id = store.start_run("run")
            posting = JobPosting(
                "2", "AI Engineer", "California", "https://example.com/2"
            )
            store.store_postings("example", [posting], run_id, baseline=False)
            signature = make_signature("example", "2")

            store.transition_processing("example", signature, "detail_pending")
            store.transition_processing(
                "example", signature, "detail_failed", error="temporary timeout"
            )
            self.assertIn(signature, store.retryable_signatures("example"))
            failed = store.get_processing_state("example", signature)
            self.assertEqual(failed["detail_attempts"], 1)
            self.assertEqual(failed["last_error"], "temporary timeout")

            store.transition_processing("example", signature, "detail_pending")
            store.transition_processing(
                "example",
                signature,
                "llm_pending",
                description="Full job description",
            )
            store.transition_processing("example", signature, "llm_pending")
            store.transition_processing(
                "example", signature, "llm_failed", error="API unavailable"
            )
            self.assertIn(signature, store.retryable_signatures("example"))

            store.transition_processing("example", signature, "llm_pending")
            store.transition_processing(
                "example",
                signature,
                "review_queue",
                model="gpt-6-luna",
                prompt_version="v1",
                evaluation={"relevance_score": 82},
            )
            completed = store.get_processing_state("example", signature)
            self.assertEqual(completed["detail_attempts"], 2)
            self.assertEqual(completed["llm_attempts"], 2)
            self.assertEqual(completed["status"], "review_queue")
            self.assertEqual(
                json.loads(completed["evaluation_json"])["relevance_score"], 82
            )
            self.assertNotIn(signature, store.retryable_signatures("example"))
            store.close()

    def test_processing_rejects_invalid_transition(self):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / "test.db")
            company = {
                "company_key": "example",
                "company": "Example",
                "connector": "greenhouse",
                "connector_key": "example",
                "careers_url": "https://example.com/careers",
            }
            store.sync_companies([company])
            run_id = store.start_run("run")
            posting = JobPosting(
                "2", "AI Engineer", "California", "https://example.com/2"
            )
            store.store_postings("example", [posting], run_id, baseline=False)
            signature = make_signature("example", "2")

            with self.assertRaisesRegex(ValueError, "Invalid processing transition"):
                store.transition_processing("example", signature, "review_queue")
            store.close()

    def test_signature_is_company_scoped(self):
        self.assertNotEqual(make_signature("a", "123"), make_signature("b", "123"))


class DailyProcessingTests(unittest.TestCase):
    @patch("job_finder.evaluate_job", return_value=evaluation_result())
    def test_llm_pending_job_reaches_review_queue(self, evaluator):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / "test.db")
            company = {
                "company_key": "example",
                "company": "Example",
                "connector": "greenhouse",
                "connector_key": "example",
                "careers_url": "https://example.com/careers",
            }
            store.sync_companies([company])
            run_id = store.start_run("run")
            posting = JobPosting(
                "2",
                "Data Engineer",
                "New York, NY",
                "https://example.com/2",
                employment_type="Full-time",
                location_type="Hybrid",
            )
            store.store_postings("example", [posting], run_id, baseline=False)
            signature = make_signature("example", "2")
            store.transition_processing("example", signature, "detail_pending")
            store.transition_processing(
                "example",
                signature,
                "llm_pending",
                description="Build reliable data systems.",
            )

            counts = evaluate_retryable_postings(
                store,
                company,
                [posting],
                POLICY,
                PROFILE,
                prompt_directory=ROOT / "config",
                config=EvaluationConfig(api_key="test-key", model="test-model"),
            )

            self.assertEqual(counts, {"llm_succeeded": 1, "llm_failed": 0})
            completed = store.get_processing_state("example", signature)
            self.assertEqual(completed["status"], "review_queue")
            self.assertEqual(completed["llm_attempts"], 1)
            self.assertEqual(completed["model"], "test-model")
            self.assertEqual(
                json.loads(completed["evaluation_json"])["relevance_score"], 88
            )
            evaluator.assert_called_once()
            store.close()

    @patch("job_finder.evaluate_job", side_effect=EvaluationError("temporary API error"))
    def test_llm_failure_stays_retryable(self, _evaluator):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / "test.db")
            company = {
                "company_key": "example",
                "company": "Example",
                "connector": "greenhouse",
                "connector_key": "example",
                "careers_url": "https://example.com/careers",
            }
            store.sync_companies([company])
            run_id = store.start_run("run")
            posting = JobPosting(
                "2", "Data Engineer", "California", "https://example.com/2"
            )
            store.store_postings("example", [posting], run_id, baseline=False)
            signature = make_signature("example", "2")
            store.transition_processing("example", signature, "detail_pending")
            store.transition_processing(
                "example", signature, "llm_pending", description="Description"
            )

            counts = evaluate_retryable_postings(
                store,
                company,
                [posting],
                POLICY,
                PROFILE,
                prompt_directory=ROOT / "config",
                config=EvaluationConfig(api_key="test-key"),
            )

            self.assertEqual(counts, {"llm_succeeded": 0, "llm_failed": 1})
            failed = store.get_processing_state("example", signature)
            self.assertEqual(failed["status"], "llm_failed")
            self.assertIn(signature, store.retryable_signatures("example"))
            store.close()

    @patch("job_finder.fetch_job_detail")
    def test_filters_fetches_details_and_retries_failures(self, fetch_detail):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / "test.db")
            company = {
                "company_key": "example",
                "company": "Example",
                "connector": "greenhouse",
                "connector_key": "example",
                "careers_url": "https://example.com/careers",
            }
            store.sync_companies([company])
            run_id = store.start_run("run")
            postings = [
                JobPosting("1", "Data Engineer", "California", "https://example.com/1", employment_type="Full-time"),
                JobPosting("2", "Sales Executive", "California", "https://example.com/2", employment_type="Full-time"),
                JobPosting("3", "Data Engineer", "Texas", "https://example.com/3", employment_type="Full-time"),
                JobPosting("4", "Data Engineer", "New York, NY", "https://example.com/4", employment_type="Contract"),
                JobPosting("5", "AI Engineer", "New York, NY", "https://example.com/5"),
            ]
            store.store_postings("example", postings, run_id, baseline=False)

            def first_attempt(_connector, _connector_key, posting):
                if posting.source_job_id == "5":
                    raise ConnectorError("temporary detail failure")
                return JobDetail(
                    "<h2>Role</h2><p>Build&nbsp;data pipelines.</p>",
                    employment_type="Full-time",
                )

            fetch_detail.side_effect = first_attempt
            counts = process_retryable_postings(store, company, postings, POLICY)

            self.assertEqual(
                counts,
                {
                    "prefilter_passed": 2,
                    "prefilter_rejected": 3,
                    "detail_succeeded": 1,
                    "detail_failed": 1,
                },
            )
            successful = store.get_processing_state(
                "example", make_signature("example", "1")
            )
            self.assertEqual(successful["status"], "llm_pending")
            self.assertEqual(successful["description"], "Role\nBuild data pipelines.")
            self.assertEqual(successful["detail_attempts"], 1)
            self.assertEqual(successful["llm_attempts"], 0)
            failed_signature = make_signature("example", "5")
            self.assertEqual(
                store.get_processing_state("example", failed_signature)["status"],
                "detail_failed",
            )

            fetch_detail.side_effect = None
            fetch_detail.return_value = JobDetail(
                "Retry succeeded.", employment_type="Full-time"
            )
            retry_counts = process_retryable_postings(store, company, postings, POLICY)

            self.assertEqual(retry_counts["detail_succeeded"], 1)
            retried = store.get_processing_state("example", failed_signature)
            self.assertEqual(retried["status"], "llm_pending")
            self.assertEqual(retried["detail_attempts"], 2)
            self.assertEqual(retried["llm_attempts"], 0)
            store.close()

    @patch("job_finder.fetch_job_detail")
    def test_marks_removed_detail_failure_terminal(self, fetch_detail):
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / "test.db")
            company = {
                "company_key": "example",
                "company": "Example",
                "connector": "greenhouse",
                "connector_key": "example",
                "careers_url": "https://example.com/careers",
            }
            store.sync_companies([company])
            run_id = store.start_run("run")
            posting = JobPosting(
                "gone", "Data Engineer", "New York, NY", "https://example.com/gone"
            )
            store.store_postings("example", [posting], run_id, baseline=False)
            signature = make_signature("example", "gone")
            store.transition_processing("example", signature, "detail_pending")
            store.transition_processing(
                "example", signature, "detail_failed", error="temporary failure"
            )

            counts = process_retryable_postings(store, company, [], POLICY)

            self.assertEqual(counts["prefilter_rejected"], 1)
            state = store.get_processing_state("example", signature)
            self.assertEqual(state["status"], "prefilter_rejected")
            self.assertIn("no longer present", state["filter_reason"])
            fetch_detail.assert_not_called()
            store.close()


class WorkbookTests(unittest.TestCase):
    def test_append_run_rows_sanitizes_external_error_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "jobs.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Run Summary"
            sheet.append([])
            sheet.append(["Run summary"])
            sheet.append(
                [
                    "Run ID", "Started At", "Completed At", "Mode", "Company",
                    "Status", "Jobs Found", "New Jobs", "Prefilter Passed",
                    "Prefilter Rejected", "Detail Succeeded", "Detail Failed",
                    "LLM Succeeded", "LLM Failed", "Review Queue Added", "Error",
                ]
            )
            sheet.append([None] * 16)
            sheet.add_table(Table(displayName="RunSummaryTable", ref="A3:P4"))
            workbook.save(workbook_path)
            workbook.close()
            append_run_rows(
                workbook_path,
                [
                    {
                        "run_id": 7,
                        "started_at": "2026-10-15T12:00:00+00:00",
                        "completed_at": "2026-10-15T12:01:00+00:00",
                        "mode": "run",
                        "company": "Meta",
                        "status": "failed",
                        "jobs_found": 0,
                        "new_jobs": 0,
                        "prefilter_passed": 0,
                        "prefilter_rejected": 0,
                        "detail_succeeded": 0,
                        "detail_failed": 0,
                        "llm_succeeded": 0,
                        "llm_failed": 0,
                        "review_queue_added": 0,
                        "error": "browser failed\x1b[2m\x07 on navigation",
                    }
                ],
            )
            workbook = load_workbook(workbook_path, data_only=False)
            sheet = workbook["Run Summary"]
            self.assertEqual(sheet.cell(4, 16).value, "browser failed on navigation")
            workbook.close()

    def test_review_queue_append_is_idempotent_by_signature_and_prompt(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Review Queue"
            sheet.append([])
            sheet.append(["Review queue"])
            headers = [
                "Date Found", "Company", "Title", "Location", "Work Arrangement",
                "Employment Type", "Experience Required", "Experience Evidence",
                "Sponsorship Status", "Sponsorship Evidence", "Relevance Score",
                "Model Recommendation", "Decision Reason", "Matching Evidence", "Gaps",
                "Job URL", "Source Job ID", "Signature", "Model", "Prompt Version",
                "User Verdict", "User Notes",
            ]
            sheet.append(headers)
            sheet.append([None] * len(headers))
            sheet.add_table(Table(displayName="ReviewQueueTable", ref="A3:V4"))
            workbook.save(path)

            evaluation = {
                "date_found": "2026-10-14T12:00:00Z",
                "company": "Example",
                "title": "Data Engineer",
                "location": "New York, NY",
                "job_url": "https://example.com/jobs/1",
                "source_job_id": "1",
                "signature": "signature-1",
                "prompt_version": "v1",
                "matching_evidence": ["Python", "SQL"],
            }
            self.assertEqual(append_review_queue_jobs(path, [evaluation]), 1)
            self.assertEqual(append_review_queue_jobs(path, [evaluation]), 0)
            self.assertTrue(review_queue_contains(path, "signature-1", "v1"))
            self.assertFalse(review_queue_contains(path, "signature-1", "v2"))

            workbook = load_workbook(path)
            sheet = workbook["Review Queue"]
            self.assertEqual(sheet["R4"].value, "signature-1")
            self.assertEqual(sheet["N4"].value, "Python; SQL")
            self.assertEqual(sheet.tables["ReviewQueueTable"].ref, "A3:V4")
            workbook.close()

    def test_table_rows_remain_contiguous_with_distant_validation(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([])
        sheet.append([])
        sheet.append(["A", "B"])
        sheet.append([None, None])
        sheet.add_table(Table(displayName="TestTable", ref="A3:B4"))
        validation = DataValidation(type="list", formula1='"Yes,No"')
        sheet.add_data_validation(validation)
        validation.add("B5000")

        _append_table_row(sheet, "TestTable", [1, 2])
        _append_table_row(sheet, "TestTable", [3, 4])

        self.assertEqual(sheet["A4"].value, 1)
        self.assertEqual(sheet["A5"].value, 3)
        self.assertEqual(sheet.tables["TestTable"].ref, "A3:B5")

    def test_daily_reader_skips_uninitialized_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "companies.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Companies"
            sheet.append([])
            sheet.append([])
            headers = [
                "Company Key",
                "Company",
                "Enabled",
                "Careers URL",
                "Connector",
                "Connector Key",
                "Status",
                "Notes",
            ]
            sheet.append(headers)
            sheet.append(["ready", "Ready Co", True, "https://ready.example", "ashby", "ready", "Ready", ""])
            sheet.append(["", "New Co", True, "https://new.example", "", "", "", ""])
            sheet.append(["pending", "Pending Co", True, "https://pending.example", "pending", "", "Pending connector", ""])
            sheet.add_table(Table(displayName="CompaniesTable", ref="A3:H6"))
            workbook.save(path)

            companies = read_companies(path)
            self.assertEqual([company["company_key"] for company in companies], ["ready"])


class ConnectorDiscoveryTests(unittest.TestCase):
    def test_extracts_common_ats_connectors(self):
        page = """
        https:\\/\\/jobs.ashbyhq.com\\/example-ai\\/abc
        https://job-boards.greenhouse.io/example/jobs/1
        https://jobs.lever.co/example-inc/abc
        https://example.wd5.myworkdayjobs.com/en-US/ExternalCareerSite/job/123
        """
        candidates = {
            (candidate.connector, candidate.connector_key)
            for candidate in extract_connector_candidates(page)
        }
        self.assertIn(("ashby", "example-ai"), candidates)
        self.assertIn(("greenhouse", "example"), candidates)
        self.assertIn(("lever", "example-inc"), candidates)
        self.assertIn(
            ("workday", "https://example.wd5.myworkdayjobs.com/ExternalCareerSite"),
            candidates,
        )

    def test_company_key_normalization(self):
        self.assertEqual(company_key_from_name("Acme & Co."), "acme_co")

    def test_greenhouse_detail_uses_description_from_board_fetch(self):
        posting = JobPosting(
            source_job_id="123",
            title="Data Engineer",
            location="New York, NY",
            job_url="https://example.com/jobs/123",
            description="  <p>Build reliable data systems.</p>  ",
            raw={"id": 123},
        )

        detail = fetch_job_detail("greenhouse", "scaleai", posting)

        self.assertIsInstance(detail, JobDetail)
        self.assertEqual(detail.description, "<p>Build reliable data systems.</p>")
        self.assertEqual(detail.raw, {"id": 123})

    def test_greenhouse_detail_rejects_missing_description(self):
        posting = JobPosting("123", "Data Engineer", "New York, NY", "https://example.com/123")

        with self.assertRaisesRegex(ConnectorError, "missing its description"):
            fetch_job_detail("greenhouse", "scaleai", posting)

    @patch("connectors._get_json")
    def test_lever_connector(self, get_json):
        get_json.return_value = [
            {
                "id": "abc",
                "text": "Data Engineer",
                "categories": {
                    "location": "New York, NY",
                    "department": "Engineering",
                    "commitment": "Full-time",
                },
                "hostedUrl": "https://jobs.lever.co/example/abc",
                "descriptionPlain": "Build data systems.",
                "createdAt": 123,
            }
        ]
        postings = fetch_lever("example")
        self.assertEqual(postings[0].source_job_id, "abc")
        self.assertEqual(postings[0].title, "Data Engineer")

    @patch("connectors.subprocess.run")
    @patch("connectors.shutil.which", return_value=r"C:\Program Files\nodejs\node.exe")
    def test_tesla_connector_maps_board_fields(self, _which, run):
        run.return_value = subprocess.CompletedProcess(
            args=[],
            returncode=0,
            stdout=json.dumps(
                {
                    "listings": [
                        {
                            "id": "283102",
                            "t": "Electrical Commissioning Lead, Cell",
                            "dp": "11",
                            "l": "445804",
                            "y": 1,
                        }
                    ],
                    "lookup": {
                        "locations": {"445804": "Buffalo, New York"},
                        "departments": {"11": "Engineering & Information Technology"},
                        "types": {"1": "fulltime"},
                    },
                }
            ),
            stderr="",
        )

        postings = fetch_tesla()

        self.assertEqual(len(postings), 1)
        self.assertEqual(postings[0].source_job_id, "283102")
        self.assertEqual(postings[0].location, "Buffalo, New York")
        self.assertEqual(postings[0].department, "Engineering & Information Technology")
        self.assertEqual(postings[0].employment_type, "fulltime")
        self.assertEqual(postings[0].description, "")
        self.assertEqual(
            postings[0].job_url,
            "https://www.tesla.com/careers/search/job/283102",
        )

    @patch("connectors.subprocess.run")
    @patch("connectors.shutil.which", return_value=r"C:\Program Files\nodejs\node.exe")
    def test_tesla_detail_combines_structured_sections(self, _which, run):
        run.return_value = subprocess.CompletedProcess(
            args=[],
            returncode=0,
            stdout=json.dumps(
                {
                    "id": "283102",
                    "jobDescription": "Build reliable data systems.",
                    "jobResponsibilities": "Own batch and streaming pipelines.",
                    "jobRequirements": "Two years of relevant experience.",
                    "jobCompensationAndBenefits": "Benefits are provided.",
                    "timeType": "Full-time",
                }
            ),
            stderr="",
        )
        posting = JobPosting(
            "283102",
            "Data Engineer",
            "Buffalo, New York",
            "https://www.tesla.com/careers/search/job/283102",
        )

        detail = fetch_job_detail("tesla", "", posting)

        self.assertIn("Job description\nBuild reliable data systems.", detail.description)
        self.assertIn("Requirements\nTwo years of relevant experience.", detail.description)
        self.assertEqual(detail.employment_type, "Full-time")
        run.assert_called_once_with(
            [
                r"C:\Program Files\nodejs\node.exe",
                str(ROOT / "src" / "tesla_browser_fetch.mjs"),
                "detail",
                "283102",
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            check=False,
        )

    @patch("connectors.subprocess.run")
    @patch("connectors.shutil.which", return_value=r"C:\Program Files\nodejs\node.exe")
    def test_tesla_detail_rejects_missing_description(self, _which, run):
        run.return_value = subprocess.CompletedProcess(
            args=[],
            returncode=0,
            stdout=json.dumps({"id": "283102", "timeType": "Full-time"}),
            stderr="",
        )
        posting = JobPosting(
            "283102",
            "Data Engineer",
            "Buffalo, New York",
            "https://www.tesla.com/careers/search/job/283102",
        )

        with self.assertRaisesRegex(ConnectorError, "missing its description"):
            fetch_job_detail("tesla", "", posting)

    @patch("connectors.fetch_tesla", return_value=[])
    def test_tesla_is_registered_in_dispatcher(self, fetch):
        self.assertEqual(fetch_jobs("tesla", ""), [])
        fetch.assert_called_once_with()

    @patch("connectors._post_json")
    def test_apple_connector_maps_and_paginates(self, post_json):
        page_one = [
            {
                "id": "200000001",
                "positionId": "200000001",
                "postingTitle": "Data Engineer",
                "transformedPostingTitle": "data-engineer",
                "locations": [{"name": "New York City, New York, United States"}],
                "team": {"teamName": "Machine Learning and AI"},
                "postDateInGMT": "2026-10-14T12:00:00Z",
                "homeOffice": False,
                "jobSummary": "Not retained during board collection.",
            }
        ]
        page_two = [
            {
                "id": "200000002",
                "positionId": "200000002",
                "postingTitle": "AI Engineer",
                "locations": [{"name": "Cupertino, California, United States"}],
                "team": {"teamName": "Software and Services"},
                "postingDate": "Oct 14, 2026",
                "homeOffice": True,
            }
        ]

        def response(_url, payload):
            listings = page_one if payload["page"] == 1 else page_two
            return {"res": {"totalRecords": 21, "searchResults": listings}}

        post_json.side_effect = response
        postings = fetch_apple("postLocation-state953,postLocation-state985")

        self.assertEqual([posting.source_job_id for posting in postings], ["200000001", "200000002"])
        self.assertEqual(postings[0].location, "New York City, New York, United States")
        self.assertEqual(postings[0].department, "Machine Learning and AI")
        self.assertEqual(postings[0].description, "")
        self.assertNotIn("jobSummary", postings[0].raw)
        self.assertEqual(postings[1].location_type, "Remote")
        self.assertEqual(
            postings[0].job_url,
            "https://jobs.apple.com/en-us/details/200000001/data-engineer",
        )

    @patch("connectors.fetch_apple", return_value=[])
    def test_apple_is_registered_in_dispatcher(self, fetch):
        self.assertEqual(fetch_jobs("apple", "postLocation-state953"), [])
        fetch.assert_called_once_with("postLocation-state953")

    @patch("connectors.subprocess.run")
    @patch("connectors.shutil.which", return_value=r"C:\Program Files\nodejs\node.exe")
    def test_browser_detail_connectors_use_company_selectors(self, _which, run):
        description = "Full job description. " * 150
        run.return_value = subprocess.CompletedProcess(
            args=[],
            returncode=0,
            stdout=json.dumps(
                {
                    "title": "Job details",
                    "url": "https://example.com/job/1",
                    "description": description,
                }
            ),
            stderr="",
        )
        cases = [
            ("apple", "#jobdetails-wrapper", "500"),
            ("adobe", ".job-desc", "1200"),
            ("meta", "body", "2000"),
        ]

        for connector, selector, minimum_length in cases:
            with self.subTest(connector=connector):
                posting = JobPosting(
                    source_job_id="1",
                    title="Data Engineer",
                    location="California",
                    job_url=f"https://example.com/{connector}/1",
                    employment_type="Full time",
                    location_type="Hybrid",
                )
                detail = fetch_job_detail(connector, "", posting)
                self.assertEqual(detail.description, description.strip())
                self.assertEqual(detail.employment_type, "Full time")
                self.assertEqual(detail.location_type, "Hybrid")
                args = run.call_args.args[0]
                self.assertEqual(args[-2:], [selector, minimum_length])

    @patch("connectors._get_json")
    def test_microsoft_connector_maps_and_deduplicates_locations(self, get_json):
        position = {
            "id": 12345,
            "displayJobId": "200000001",
            "name": "Data Engineer",
            "locations": ["United States, New York, New York"],
            "department": "Software Engineering",
            "workLocationOption": "hybrid",
            "postedTs": 100,
            "positionUrl": "/careers/job/12345",
        }
        get_json.return_value = {
            "status": 200,
            "data": {"count": 1, "positions": [position]},
        }

        postings = fetch_microsoft("California|New York")

        self.assertEqual(len(postings), 1)
        self.assertEqual(postings[0].source_job_id, "200000001")
        self.assertEqual(postings[0].location_type, "hybrid")
        self.assertEqual(
            postings[0].job_url,
            "https://apply.careers.microsoft.com/careers/job/12345",
        )
        self.assertEqual(get_json.call_count, 2)
        first_kwargs = get_json.call_args_list[0].kwargs
        second_kwargs = get_json.call_args_list[1].kwargs
        self.assertIn("Mozilla/5.0", first_kwargs["user_agent"])
        self.assertIs(first_kwargs["opener"], second_kwargs["opener"])

    @patch("connectors.time.sleep")
    @patch("connectors._fetch_microsoft_page")
    def test_microsoft_connector_merges_retried_page_boundaries(self, fetch_page, _sleep):
        positions = [
            {
                "id": index,
                "displayJobId": f"JOB{index}",
                "name": f"Role {index}",
                "locations": ["California"],
            }
            for index in range(11)
        ]
        fetch_page.side_effect = [
            (11, positions[:10]),
            (11, [positions[9]]),
            (11, positions[:10]),
            (11, [positions[10]]),
        ]

        postings = fetch_microsoft("California")

        self.assertEqual(len(postings), 11)
        self.assertEqual({posting.source_job_id for posting in postings}, {f"JOB{i}" for i in range(11)})

    @patch("connectors._get_text")
    def test_microsoft_detail_reads_jobposting_json_ld(self, get_text):
        get_text.return_value = """
        <html><head><script type="application/ld+json">
        {
          "@context": "https://schema.org",
          "@type": "JobPosting",
          "description": "<p>Build reliable data services.</p>",
          "employmentType": "FULL_TIME"
        }
        </script></head></html>
        """
        posting = JobPosting(
            "200000001",
            "Data Engineer",
            "California",
            "https://apply.careers.microsoft.com/careers/job/12345",
            location_type="Hybrid",
        )

        detail = fetch_job_detail("microsoft", "", posting)

        self.assertEqual(detail.description, "<p>Build reliable data services.</p>")
        self.assertEqual(detail.employment_type, "FULL_TIME")
        self.assertEqual(detail.location_type, "Hybrid")

    @patch("connectors.subprocess.run")
    @patch("connectors.shutil.which", return_value=r"C:\Program Files\nodejs\node.exe")
    def test_meta_connector_maps_browser_results(self, _which, run):
        run.return_value = subprocess.CompletedProcess(
            args=[],
            returncode=0,
            stdout=json.dumps(
                {
                    "job_count": 1,
                    "listings": [
                        {
                            "id": "123",
                            "title": "Product Manager, Data",
                            "locations": ["New York, NY"],
                            "teams": ["Product Management"],
                        }
                    ],
                }
            ),
            stderr="",
        )

        postings = fetch_meta()

        self.assertEqual(postings[0].source_job_id, "123")
        self.assertEqual(postings[0].location, "New York, NY")
        self.assertEqual(postings[0].department, "Product Management")

    @patch("connectors._post_json")
    def test_adobe_connector_reads_stably_sorted_pages(self, post_json):
        first = {
            "status": 200,
            "totalHits": 11,
            "data": {
                "jobs": [
                    {
                        "reqId": f"R{index}",
                        "title": "AI Engineer" if index == 1 else f"Role {index}",
                        "location": "San Jose, California, United States",
                        "category": "Engineering and Product",
                        "type": "Full time",
                        "descriptionTeaser": "Not stored.",
                    }
                    for index in range(1, 11)
                ]
            },
        }
        second = {
            "status": 200,
            "totalHits": 11,
            "data": {
                "jobs": [
                    {
                        "reqId": "R11",
                        "title": "Data Engineer",
                        "multi_location": ["New York, New York, United States"],
                        "multi_category": ["Engineering and Product"],
                        "type": "Full time",
                    }
                ]
            },
        }
        post_json.side_effect = [
            {"refineSearch": first},
            {"refineSearch": second},
        ]

        postings = fetch_adobe("ADOBUS")

        self.assertEqual(
            [posting.source_job_id for posting in postings],
            [f"R{index}" for index in range(1, 12)],
        )
        self.assertEqual(postings[-1].location, "New York, New York, United States")
        self.assertNotIn("descriptionTeaser", postings[0].raw)

    @patch("connectors._get_json")
    def test_salesforce_connector_maps_static_feed(self, get_json):
        get_json.return_value = {
            "Count": 1,
            "Report_Entry": [
                {
                    "Job_Requisition_Ref_ID": "JR1",
                    "Job_Posting_Title": "Data Engineer",
                    "Job_Requisition_Primary_Location": "New York - New York",
                    "Job_Requisition_Additional_Locations": "Massachusetts - Boston",
                    "Job_Family_Group": "Data",
                    "Time_Type": "Full time",
                    "External_Job_Posting_Start_Date": "2026-10-14",
                    "Job_Description": "Not stored during board collection.",
                }
            ],
        }

        postings = fetch_salesforce("prod")

        self.assertEqual(postings[0].source_job_id, "JR1")
        self.assertEqual(
            postings[0].location,
            "New York - New York; Massachusetts - Boston",
        )
        self.assertEqual(postings[0].description, "")
        self.assertNotIn("Job_Description", postings[0].raw)

        detail = fetch_job_detail("salesforce", "prod", postings[0])

        self.assertEqual(detail.description, "Not stored during board collection.")
        self.assertEqual(detail.employment_type, "Full time")
        self.assertIn("Job_Description", detail.raw)
        get_json.assert_called_once()

    @patch("connectors._get_json")
    def test_salesforce_detail_can_load_feed_without_board_cache(self, get_json):
        get_json.return_value = {
            "Count": 1,
            "Report_Entry": [
                {
                    "Job_Requisition_Ref_ID": "JR2",
                    "Job_Posting_Title": "AI Engineer",
                    "Time_Type": "Full time",
                    "Job_Description": "Build machine-learning infrastructure.",
                }
            ],
        }
        posting = JobPosting(
            "JR2",
            "AI Engineer",
            "California",
            "https://careers.salesforce.com/en/jobs/jr2/",
        )

        detail = fetch_job_detail("salesforce", "uat", posting)

        self.assertEqual(detail.description, "Build machine-learning infrastructure.")
        get_json.assert_called_once()

    @patch("connectors._get_json")
    def test_salesforce_detail_falls_back_to_workday(self, get_json):
        get_json.return_value = {
            "jobPostingInfo": {
                "jobDescription": "<p>Build reliable analytics systems.</p>"
            }
        }
        posting = JobPosting(
            "JR3",
            "Data Engineer",
            "New York",
            "https://careers.salesforce.com/en/jobs/jr3/",
            employment_type="Full time",
            raw={
                "External_Job_Posting_Site": (
                    "https://salesforce.wd12.myworkdayjobs.com/"
                    "External_Career_Site/job/New-York/Data-Engineer_JR3"
                )
            },
        )
        _SALESFORCE_DETAIL_CACHE["prod"] = {
            "JR3": {
                "Job_Requisition_Ref_ID": "JR3",
                "Time_Type": "Full time",
                "External_Job_Posting_Site": posting.raw[
                    "External_Job_Posting_Site"
                ],
            }
        }

        detail = fetch_job_detail("salesforce", "prod", posting)

        self.assertEqual(detail.description, "<p>Build reliable analytics systems.</p>")
        url = get_json.call_args.args[0]
        self.assertEqual(
            url,
            "https://salesforce.wd12.myworkdayjobs.com/wday/cxs/salesforce/"
            "External_Career_Site/job/New-York/Data-Engineer_JR3",
        )
        self.assertIn("Mozilla/5.0", get_json.call_args.kwargs["user_agent"])

    @patch("connectors._get_json")
    def test_smartrecruiters_connector_maps_and_paginates(self, get_json):
        first_job = {
            "id": "101",
            "name": "Data Engineer",
            "releasedDate": "2026-10-14T12:00:00Z",
            "location": {
                "fullLocation": "New York, NY, United States",
                "remote": False,
                "hybrid": True,
            },
            "department": {"label": "Engineering"},
            "typeOfEmployment": {"label": "Full-time"},
        }
        second_job = {
            "id": "102",
            "name": "AI Engineer",
            "location": {
                "city": "Sunnyvale",
                "region": "CA",
                "country": "us",
                "remote": False,
                "hybrid": False,
            },
            "function": {"label": "Engineering"},
        }
        get_json.side_effect = [
            {"totalFound": 101, "content": [first_job]},
            {"totalFound": 101, "content": [second_job]},
        ]

        postings = fetch_smartrecruiters("LinkedIn3")

        self.assertEqual(
            [posting.source_job_id for posting in postings], ["101", "102"]
        )
        self.assertEqual(postings[0].location_type, "Hybrid")
        self.assertEqual(postings[1].location, "Sunnyvale, CA, us")
        self.assertEqual(
            postings[0].job_url,
            "https://jobs.smartrecruiters.com/LinkedIn3/101",
        )
        self.assertEqual(postings[0].description, "")

    @patch("connectors.fetch_smartrecruiters", return_value=[])
    def test_smartrecruiters_is_registered_in_dispatcher(self, fetch):
        self.assertEqual(fetch_jobs("smartrecruiters", "LinkedIn3"), [])
        fetch.assert_called_once_with("LinkedIn3")

    @patch("connectors._get_json")
    def test_smartrecruiters_detail_combines_sections(self, get_json):
        get_json.return_value = {
            "jobAd": {
                "sections": {
                    "jobDescription": {
                        "title": "Job Description",
                        "text": "Build data products.",
                    },
                    "qualifications": {
                        "title": "Qualifications",
                        "text": "Two years of experience.",
                    },
                    "additionalInformation": {
                        "title": "Additional Information",
                        "text": "Hybrid role.",
                    },
                }
            },
            "typeOfEmployment": {"label": "Full-time"},
            "location": {"hybrid": True, "remote": False},
        }
        posting = JobPosting(
            "101",
            "Data Engineer",
            "Sunnyvale, CA",
            "https://jobs.smartrecruiters.com/LinkedIn3/101",
        )

        detail = fetch_job_detail("smartrecruiters", "LinkedIn3", posting)

        self.assertIn("Job Description\nBuild data products.", detail.description)
        self.assertIn("Qualifications\nTwo years of experience.", detail.description)
        self.assertEqual(detail.employment_type, "Full-time")
        self.assertEqual(detail.location_type, "Hybrid")
        get_json.assert_called_once_with(
            "https://api.smartrecruiters.com/v1/companies/LinkedIn3/postings/101"
        )

    @patch("connectors.fetch_adobe", return_value=[])
    @patch("connectors.fetch_meta", return_value=[])
    @patch("connectors.fetch_microsoft", return_value=[])
    @patch("connectors.fetch_salesforce", return_value=[])
    def test_remaining_connectors_are_registered(
        self, salesforce, microsoft, meta, adobe
    ):
        self.assertEqual(fetch_jobs("adobe", "ADOBUS"), [])
        self.assertEqual(fetch_jobs("meta", "jobsearch"), [])
        self.assertEqual(fetch_jobs("microsoft", "California"), [])
        self.assertEqual(fetch_jobs("salesforce", "prod"), [])
        adobe.assert_called_once_with("ADOBUS")
        meta.assert_called_once_with()
        microsoft.assert_called_once_with("California")
        salesforce.assert_called_once_with("prod")


class InitializerTests(unittest.TestCase):
    @staticmethod
    def _create_workbook(path: Path) -> None:
        workbook = Workbook()
        companies = workbook.active
        companies.title = "Companies"
        companies.append([])
        companies.append(["Companies"])
        companies.append(
            [
                "Company Key",
                "Company",
                "Enabled",
                "Careers URL",
                "Connector",
                "Connector Key",
                "Status",
                "Notes",
            ]
        )
        companies.append(["", "Example AI", None, "https://example.com/careers", "", "", "", ""])
        companies.add_table(Table(displayName="CompaniesTable", ref="A3:H4"))

        summary = workbook.create_sheet("Run Summary")
        summary.append([])
        summary.append(["Run summary"])
        summary.append(
            [
                "Run ID",
                "Run Started",
                "Run Completed",
                "Mode",
                "Company",
                "Status",
                "Jobs Found",
                "New Jobs",
                "Prefilter Passed",
                "Prefilter Rejected",
                "Detail Succeeded",
                "Detail Failed",
                "LLM Succeeded",
                "LLM Failed",
                "Review Queue Added",
                "Error",
            ]
        )
        summary.append([None] * 16)
        summary.add_table(Table(displayName="RunSummaryTable", ref="A3:P4"))
        workbook.save(path)

    def test_initializes_only_once(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "jobs.xlsx"
            database_path = Path(directory) / "jobs.db"
            self._create_workbook(workbook_path)
            result = DiscoveryResult(
                connector="ashby",
                connector_key="example",
                careers_url="https://example.com/careers",
                status="verified",
                note="Verified test feed.",
                postings=(
                    JobPosting("1", "Data Engineer", "New York, NY", "https://example.com/1"),
                ),
            )

            with patch("initialize_companies.resolve_company", return_value=result) as resolver:
                self.assertEqual(initialize_companies(workbook_path, database_path), 0)
                resolver.assert_called_once()

            workbook = load_workbook(workbook_path)
            companies = workbook["Companies"]
            self.assertEqual(companies["A4"].value, "example_ai")
            self.assertEqual(companies["E4"].value, "ashby")
            self.assertEqual(companies["F4"].value, "example")
            self.assertEqual(companies["G4"].value, "Initialized")
            workbook.close()

            with patch("initialize_companies.resolve_company") as resolver:
                self.assertEqual(initialize_companies(workbook_path, database_path), 0)
                resolver.assert_not_called()

            store = Store(database_path)
            self.assertTrue(store.is_initialized("example_ai"))
            count = store.connection.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]
            self.assertEqual(count, 1)
            store.close()


if __name__ == "__main__":
    unittest.main()
