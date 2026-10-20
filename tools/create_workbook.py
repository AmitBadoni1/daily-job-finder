from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "outputs" / "job_finder.xlsx"

NAVY = "17365D"
WHITE = "FFFFFF"


def _prepare_sheet(sheet, title: str, headers: list[str], table_name: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A2"] = title
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
    sheet.append([])
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(3, column, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[3].height = 30
    sheet.append([None] * len(headers))
    end_column = sheet.cell(3, len(headers)).column_letter
    table = Table(displayName=table_name, ref=f"A3:{end_column}4")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A4"


def create_workbook(path: Path, *, force: bool = False) -> None:
    if path.exists() and not force:
        print(f"Workbook already exists: {path}")
        return

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    jobs = workbook.create_sheet("Relevant Jobs")
    _prepare_sheet(
        jobs,
        "Relevant jobs",
        [
            "Discovered Date",
            "Company",
            "Job Title",
            "Location",
            "Work Arrangement",
            "Experience Requirement",
            "Job URL",
            "Relevance Score",
            "Decision Reason",
            "Matching Evidence",
            "Gaps",
            "Application Status",
            "Notes",
            "Job ID",
        ],
        "RelevantJobsTable",
    )
    for column, width in {
        "A": 18,
        "B": 18,
        "C": 34,
        "D": 24,
        "E": 20,
        "F": 24,
        "G": 42,
        "H": 16,
        "I": 34,
        "J": 34,
        "K": 28,
        "L": 18,
        "M": 28,
        "N": 18,
    }.items():
        jobs.column_dimensions[column].width = width
    application_status = DataValidation(
        type="list",
        formula1='"Not started,Applied,Interviewing,Rejected,Closed"',
    )
    jobs.add_data_validation(application_status)
    application_status.add("L4:L5000")

    review = workbook.create_sheet("Review Queue")
    _prepare_sheet(
        review,
        "Review queue",
        [
            "Date Found",
            "Company",
            "Title",
            "Location",
            "Work Arrangement",
            "Employment Type",
            "Experience Required",
            "Experience Evidence",
            "Sponsorship Status",
            "Sponsorship Evidence",
            "Relevance Score",
            "Model Recommendation",
            "Decision Reason",
            "Matching Evidence",
            "Gaps",
            "Job URL",
            "Source Job ID",
            "Signature",
            "Model",
            "Prompt Version",
            "User Verdict",
            "User Notes",
        ],
        "ReviewQueueTable",
    )
    for column in range(1, 23):
        review.column_dimensions[review.cell(3, column).column_letter].width = 22
    review.column_dimensions["C"].width = 34
    review.column_dimensions["M"].width = 34
    review.column_dimensions["N"].width = 34
    review.column_dimensions["P"].width = 42
    review.column_dimensions["V"].width = 34
    verdict = DataValidation(type="list", formula1='"Accept,Reject,Unsure"')
    review.add_data_validation(verdict)
    verdict.add("U4:U5000")

    summary = workbook.create_sheet("Run Summary")
    _prepare_sheet(
        summary,
        "Run summary",
        [
            "Run ID",
            "Run Started",
            "Run Completed",
            "Mode",
            "Company",
            "Status",
            "Jobs Found",
            "New Jobs",
            "Prefilter Passed",
            "Prefilter Rejected",
            "Detail Succeeded",
            "Detail Failed",
            "LLM Succeeded",
            "LLM Failed",
            "Review Queue Added",
            "Error",
        ],
        "RunSummaryTable",
    )
    for column in range(1, 17):
        summary.column_dimensions[summary.cell(3, column).column_letter].width = 16
    summary.column_dimensions["B"].width = 22
    summary.column_dimensions["C"].width = 22
    summary.column_dimensions["P"].width = 55

    companies = workbook.create_sheet("Companies")
    _prepare_sheet(
        companies,
        "Companies",
        [
            "Company Key",
            "Company",
            "Enabled",
            "Careers URL",
            "Connector",
            "Connector Key",
            "Status",
            "Notes",
        ],
        "CompaniesTable",
    )
    included_companies = [
        ["tesla", "Tesla", False, "https://www.tesla.com/careers/search/", "tesla", "tesla", "Ready", "Included connector"],
        ["apple", "Apple", False, "https://jobs.apple.com/en-us/search", "apple", "postLocation-state953,postLocation-state985,postLocation-BSTNMETRO", "Ready", "Included connector; location IDs match the example policy"],
        ["microsoft", "Microsoft", False, "https://jobs.careers.microsoft.com/global/en/search", "microsoft", "California|New York|Boston, Massachusetts", "Ready", "Included connector; searches match the example policy"],
        ["scale_ai", "Scale AI", False, "https://scale.com/careers", "greenhouse", "scaleai", "Ready", "Included Greenhouse connector"],
        ["meta", "Meta", False, "https://www.metacareers.com/jobs/", "meta", "jobsearch", "Ready", "Included connector"],
        ["adobe", "Adobe", False, "https://careers.adobe.com/us/en/search-results", "adobe", "ADOBUS", "Ready", "Included connector"],
        ["salesforce", "Salesforce", False, "https://careers.salesforce.com/en/jobs/", "salesforce", "prod", "Ready", "Included connector"],
        ["linkedin", "LinkedIn", False, "https://careers.smartrecruiters.com/LinkedIn3", "smartrecruiters", "LinkedIn3", "Ready", "Included SmartRecruiters connector"],
    ]
    companies.delete_rows(4, 1)
    for row in included_companies:
        companies.append(row)
    companies.tables["CompaniesTable"].ref = f"A3:H{3 + len(included_companies)}"
    for column, width in {
        "A": 22,
        "B": 24,
        "C": 12,
        "D": 55,
        "E": 20,
        "F": 24,
        "G": 20,
        "H": 42,
    }.items():
        companies.column_dimensions[column].width = width
    enabled = DataValidation(type="list", formula1='"TRUE,FALSE"')
    companies.add_data_validation(enabled)
    enabled.add("C4:C5000")

    rules = workbook.create_sheet("Search Rules")
    _prepare_sheet(
        rules,
        "Search rules",
        ["Rule", "Value", "Notes"],
        "SearchRulesTable",
    )
    rule_rows = [
        ["Configuration", "config/search_policy.json", "Edit the JSON file to change screening rules."],
        ["Candidate profile", "config/candidate_profile.json", "Summarize only evidence that is present in your resume."],
        ["Duplicate rule", "Company plus source job ID", "A previously seen ID is not treated as new."],
        ["Output", "Review Queue", "New jobs that pass local filters are evaluated by the model."],
    ]
    rules.delete_rows(4, 1)
    for row in rule_rows:
        rules.append(row)
    rules.tables["SearchRulesTable"].ref = f"A3:C{3 + len(rule_rows)}"
    rules.column_dimensions["A"].width = 28
    rules.column_dimensions["B"].width = 38
    rules.column_dimensions["C"].width = 70

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=4):
            for cell in row:
                cell.font = Font(name="Arial", size=10, color="1F2937")
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    workbook.save(path)
    workbook.close()
    print(f"Created workbook: {path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Create a fresh Job Finder workbook")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--force", action="store_true")
    arguments = parser.parse_args()
    create_workbook(arguments.output, force=arguments.force)


if __name__ == "__main__":
    main()
