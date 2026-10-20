from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


path = Path(__file__).resolve().parents[1] / "outputs" / "job_finder.xlsx"
workbook = load_workbook(path, data_only=False)
for sheet_name in workbook.sheetnames:
    sheet = workbook[sheet_name]
    nonempty = sum(
        1
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    print(f"{sheet_name}: max_row={sheet.max_row} max_column={sheet.max_column} nonempty={nonempty}")
    if sheet_name == "Run Summary":
        for values in sheet.iter_rows(min_row=3, values_only=True):
            if any(value is not None for value in values):
                print(values)
workbook.close()
