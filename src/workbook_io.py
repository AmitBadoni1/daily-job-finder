from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.table import Table


COMPANY_HEADERS = {
    "company_key": "Company Key",
    "company": "Company",
    "enabled": "Enabled",
    "careers_url": "Careers URL",
    "connector": "Connector",
    "connector_key": "Connector Key",
    "status": "Status",
    "notes": "Notes",
}


ANSI_ESCAPE_RE = re.compile(r"\x1b\[[0-?]*[ -/]*[@-~]")
EXCEL_CELL_CHARACTER_LIMIT = 32_767


def _header_map(sheet, row: int = 3) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in sheet[row]
        if cell.value is not None and str(cell.value).strip()
    }


def _enabled(value, *, blank_default: bool = False) -> bool:
    if value is None or not str(value).strip():
        return blank_default
    return str(value).strip().lower() in {"true", "yes", "1", "enabled"}


def read_company_rows(path: str | Path) -> list[dict]:
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Companies"]
    headers = _header_map(sheet)
    required = list(COMPANY_HEADERS.values())
    missing = [name for name in required if name not in headers]
    if missing:
        raise ValueError(f"Companies sheet is missing columns: {missing}")

    companies: list[dict] = []
    for row in range(4, sheet.max_row + 1):
        company = sheet.cell(row, headers["Company"]).value
        if company is None or not str(company).strip():
            continue
        enabled_value = sheet.cell(row, headers["Enabled"]).value
        companies.append(
            {
                "row_number": row,
                "company_key": str(sheet.cell(row, headers["Company Key"]).value or "").strip(),
                "company": str(company).strip(),
                "enabled": _enabled(enabled_value, blank_default=True),
                "careers_url": str(sheet.cell(row, headers["Careers URL"]).value or "").strip(),
                "connector": str(sheet.cell(row, headers["Connector"]).value or "").strip(),
                "connector_key": str(sheet.cell(row, headers["Connector Key"]).value or "").strip(),
                "status": str(sheet.cell(row, headers["Status"]).value or "").strip(),
                "notes": str(sheet.cell(row, headers["Notes"]).value or "").strip(),
            }
        )
    workbook.close()
    return companies


def read_companies(path: str | Path) -> list[dict]:
    companies: list[dict] = []
    for company in read_company_rows(path):
        if not company["enabled"]:
            continue
        status = company["status"].strip().lower()
        if status and status not in {"ready", "initialized"}:
            continue
        if not company["company_key"] or not company["connector"] or not company["connector_key"]:
            continue
        companies.append(company)
    return companies


def update_company_rows(path: str | Path, updates: list[dict]) -> None:
    if not updates:
        return
    workbook = load_workbook(path)
    sheet = workbook["Companies"]
    headers = _header_map(sheet)
    missing = [header for header in COMPANY_HEADERS.values() if header not in headers]
    if missing:
        workbook.close()
        raise ValueError(f"Companies sheet is missing columns: {missing}")

    highest_row = 4
    for update in updates:
        row_number = int(update["row_number"])
        highest_row = max(highest_row, row_number)
        for field, header in COMPANY_HEADERS.items():
            if field in update:
                sheet.cell(row_number, headers[header], update[field])

    table = next(
        (item for item in sheet.tables.values() if item.name == "CompaniesTable"),
        None,
    )
    if not isinstance(table, Table):
        workbook.close()
        raise ValueError("Table CompaniesTable not found on sheet Companies")
    start, end = table.ref.split(":")
    if sheet[end].row < highest_row:
        table.ref = f"{start}:{sheet.cell(highest_row, sheet[end].column).coordinate}"

    workbook.save(path)
    workbook.close()


def _append_table_row(sheet, table_name: str, values: list) -> None:
    table = next((item for item in sheet.tables.values() if item.name == table_name), None)
    if not isinstance(table, Table):
        raise ValueError(f"Table {table_name} not found on sheet {sheet.title}")
    start, end = table.ref.split(":")
    header_row = sheet[start].row
    table_end_row = sheet[end].row
    first_data_row = header_row + 1
    is_placeholder = all(sheet.cell(first_data_row, col).value is None for col in range(1, len(values) + 1))
    target_row = first_data_row if is_placeholder else table_end_row + 1
    for column, value in enumerate(values, start=1):
        sheet.cell(target_row, column, _excel_safe_value(value))
    table.ref = f"{start}:{sheet.cell(target_row, len(values)).coordinate}"


def _excel_safe_value(value):
    """Remove control sequences that cannot be stored in an XLSX text cell."""
    if not isinstance(value, str):
        return value
    value = ANSI_ESCAPE_RE.sub("", value)
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return value[:EXCEL_CELL_CHARACTER_LIMIT]


def _excel_datetime(value: str | datetime) -> datetime:
    """Return a timezone-naive datetime that Excel can store as a typed date."""
    if isinstance(value, datetime):
        parsed = value
    else:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone()
    return parsed.replace(tzinfo=None)


def append_run_rows(path: str | Path, rows: list[dict]) -> None:
    workbook = load_workbook(path)
    sheet = workbook["Run Summary"]
    for row in rows:
        _append_table_row(
            sheet,
            "RunSummaryTable",
            [
                row["run_id"],
                _excel_datetime(row["started_at"]),
                _excel_datetime(row["completed_at"]),
                row["mode"],
                row["company"],
                row["status"],
                row["jobs_found"],
                row["new_jobs"],
                row.get("prefilter_passed", row.get("title_candidates", 0)),
                row.get("prefilter_rejected", 0),
                row.get("detail_succeeded", 0),
                row.get("detail_failed", 0),
                row.get("llm_succeeded", 0),
                row.get("llm_failed", 0),
                row.get("review_queue_added", 0),
                row.get("error", ""),
            ],
        )
    workbook.save(path)
    workbook.close()


def _display_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "; ".join(str(item) for item in value if str(item).strip())
    return str(value)


def append_review_queue_jobs(path: str | Path, jobs: list[dict]) -> int:
    """Append new evaluations and return the number of rows actually added."""
    if not jobs:
        return 0
    workbook = load_workbook(path)
    sheet = workbook["Review Queue"]
    headers = _header_map(sheet)
    required = ["Signature", "Prompt Version"]
    missing = [header for header in required if header not in headers]
    if missing:
        workbook.close()
        raise ValueError(f"Review Queue sheet is missing columns: {missing}")

    existing = {
        (
            str(sheet.cell(row, headers["Signature"]).value or "").strip(),
            str(sheet.cell(row, headers["Prompt Version"]).value or "").strip(),
        )
        for row in range(4, sheet.max_row + 1)
        if sheet.cell(row, headers["Signature"]).value
    }
    added = 0
    for job in jobs:
        key = (str(job["signature"]), str(job["prompt_version"]))
        if key in existing:
            continue
        _append_table_row(
            sheet,
            "ReviewQueueTable",
            [
                _excel_datetime(job.get("date_found", datetime.now())),
                job["company"],
                job["title"],
                job["location"],
                job.get("work_arrangement", ""),
                job.get("employment_type", ""),
                job.get("experience_required", ""),
                _display_value(job.get("experience_evidence")),
                job.get("sponsorship_status", ""),
                _display_value(job.get("sponsorship_evidence")),
                job.get("relevance_score", ""),
                job.get("model_recommendation", ""),
                job.get("decision_reason", ""),
                _display_value(job.get("matching_evidence")),
                _display_value(job.get("gaps")),
                job["job_url"],
                job["source_job_id"],
                job["signature"],
                job.get("model", ""),
                job["prompt_version"],
                job.get("user_verdict", ""),
                job.get("user_notes", ""),
            ],
        )
        existing.add(key)
        added += 1
    workbook.save(path)
    workbook.close()
    return added


def review_queue_contains(
    path: str | Path, signature: str, prompt_version: str
) -> bool:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook["Review Queue"]
    headers = _header_map(sheet)
    signature_index = headers["Signature"] - 1
    prompt_index = headers["Prompt Version"] - 1
    for values in sheet.iter_rows(min_row=4, values_only=True):
        if (
            str(values[signature_index] or "").strip() == signature
            and str(values[prompt_index] or "").strip() == prompt_version
        ):
            workbook.close()
            return True
    workbook.close()
    return False


def append_relevant_jobs(path: str | Path, jobs: list[dict]) -> None:
    if not jobs:
        return
    workbook = load_workbook(path)
    sheet = workbook["Relevant Jobs"]
    for job in jobs:
        _append_table_row(
            sheet,
            "RelevantJobsTable",
            [
                datetime.now(),
                job["company"],
                job["title"],
                job["location"],
                job.get("location_type", ""),
                job.get("experience_requirement", ""),
                job["job_url"],
                job.get("relevance_score", ""),
                job.get("decision_reason", ""),
                job.get("matching_evidence", ""),
                job.get("gaps", ""),
                "Not started",
                "",
                job["source_job_id"],
            ],
        )
    workbook.save(path)
    workbook.close()
